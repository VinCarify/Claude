#!/usr/bin/env python3
"""End-to-end test for build_sheet.py.

Builds the sample workbook, lets LibreOffice (headless) calculate it, reads
the values back and compares them against an independent pure-Python
reference implementation of the KM business logic below.

    python3 test_build_sheet.py            # needs soffice (LibreOffice Calc)

Temp files go to $SHEET_TEST_TMP (default: <system tmp>/km-sheet-test).
"""
from __future__ import annotations

import calendar
import csv
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
BUILD = HERE / "build_sheet.py"
FIX_INPUT = HERE / "fixtures" / "sample_input.csv"
FIX_DEPOSITS = HERE / "fixtures" / "sample_deposits.csv"
TMP = Path(os.environ.get("SHEET_TEST_TMP")
           or Path(tempfile.gettempdir()) / "km-sheet-test")

CONTRACT = (
    "vehicle,plate,status,insurance,km_package,min_duration_months,"
    "min_duration_left,months_elapsed,monthly_price_chf,handover_date,"
    "handover_km,segment_start,min_term_end,checkpoint_date,checkpoint_type,"
    "current_km,driven_km,included_km,diff_km,extra_km_price_rp,"
    "extra_km_cost_chf,deposit_chf,open_extra_km,avg_km_month,"
    "projected_km_end,included_km_end,projected_diff_end,"
    "projected_cost_end_chf,recommended_package").split(",")
TEXT_COLS = {"vehicle", "plate", "status", "insurance", "checkpoint_type"}
TABS = ["Anleitung", "Übersicht", "web", "calc", "data_input",
        "data_deposits", "Config"]

ALLOWED_FUNCS = {
    "IF", "IFERROR", "AND", "OR", "NOT", "INDEX", "MATCH", "COUNTIF",
    "COUNTIFS", "SUMIF", "SUMIFS", "MINIFS", "MAXIFS", "SUMPRODUCT",
    "DATEDIF", "EDATE", "DATE", "LEFT", "MID", "VALUE", "ISNUMBER", "ISBLANK",
    "LEN", "TRIM", "ROUND", "MAX", "MIN", "ABS", "TODAY", "N",
    "ISERROR", "NA",  # import / MwSt-Faktor checks in Config
    "SUBSTITUTE",     # only for "250,5" -> "250.5" in deposit amounts
    "IMPORTDATA",     # only data_input!A1 / data_deposits!A1 (checked below)
}
FORBIDDEN_FUNCS = {
    "FILTER", "LET", "LAMBDA", "MAP", "REDUCE", "SCAN", "MAKEARRAY", "BYROW",
    "BYCOL", "XLOOKUP", "XMATCH", "ARRAYFORMULA", "QUERY", "AGGREGATE",
    "SORT", "SORTN", "SORTBY", "UNIQUE", "SEQUENCE", "IFS", "SWITCH",
    "TEXTJOIN", "CONCAT", "VSTACK", "HSTACK", "TOCOL", "TOROW", "CHOOSECOLS",
    "INDIRECT", "OFFSET",
}
ERRORS = {"#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!"}

# ---------------------------------------------------------------------------
# Independent reference implementation (spec, not formulas)
# ---------------------------------------------------------------------------
EPOCH = dt.date(1899, 12, 30)
ACTIVE = {"on_lease", "notice_given"}
# Fake tariff values (the real ones are company data and not in the repo)
FIX_TARIFFS = HERE / "fixtures" / "tariffs.fake.json"
_T = json.loads(FIX_TARIFFS.read_text(encoding="utf-8"))
CUTOFF = dt.date.fromisoformat(_T["cutoff"])
OLD_BRACKETS = [tuple(r) for r in _T["old"] if r[0] is not None]  # above: none
NEW_BRACKETS = [tuple(r) for r in _T["new"] if r[0] is not None]
NEW_ABOVE = _T["new"][-1][1]
PACKAGES = _T["packages"]


def serial(d: dt.date) -> int:
    return (d - EPOCH).days


def xround(x: float) -> int:
    """Spreadsheet ROUND(x, 0): half away from zero."""
    return int(Decimal(repr(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def to_num(s: str, default=None, comma=False):
    s = s.strip()
    if comma:
        s = s.replace(",", ".")
    if not s:
        return default
    try:
        return float(s)
    except ValueError:
        return default


def to_date(s: str):
    s = s.strip()
    if len(s) < 10:
        return None
    return dt.date(int(s[0:4]), int(s[5:7]), int(s[8:10]))


def datedif_months(a: dt.date, b: dt.date) -> int:
    m = (b.year - a.year) * 12 + b.month - a.month
    return m - 1 if b.day < a.day else m


def edate(d: dt.date, n: int) -> dt.date:
    y, m = divmod(d.month - 1 + n, 12)
    y += d.year
    return dt.date(y, m + 1, min(d.day, calendar.monthrange(y, m + 1)[1]))


def extra_km_price(chain_start, monthly, listing):
    if chain_start is None:
        return None
    if chain_start < CUTOFF:
        if monthly is None:
            return None
        return next((p for ub, p in OLD_BRACKETS if monthly <= ub), None)
    if listing is None:
        return None
    return next((p for ub, p in NEW_BRACKETS if listing <= ub), NEW_ABOVE)


def read_csv(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def reference(inputs: list[dict], deposits: list[dict], as_of: dt.date):
    rows = []
    for i, raw in enumerate(inputs):
        status = raw["subscription_status"].strip()
        if not status:
            continue
        rows.append({
            "sheet_row": i + 2,
            "status": status,
            "active": status in ACTIVE,
            "key": (raw["last_name"].strip().lower(),
                    raw["license_plate_number"].strip().lower()),
            "vehicle": raw["vehicle"].strip(),
            "plate": raw["license_plate_number"].strip(),
            "insurance": raw["insurance_package_type"].strip(),
            "checkpoint_type": raw["checkpoint_type"].strip(),
            "km": to_num(raw["km"], 0),
            "min_m": to_num(raw["minimum_duration_months"], 0),
            "start": to_num(raw["start_odometer"], 0),
            "adj": to_num(raw["total_km_adjustment"], 0),
            "current": to_num(raw["current_odometer"]),
            "listing": to_num(raw["s_listing"]),
            "monthly": to_num(raw["monthly_price"]),
            "handover": to_date(raw["handover_date"]),
            "cancel": to_date(raw["cancellation_date"]),
            "checkpoint": to_date(raw["checkpoint_date"]),
        })

    def chain(r):
        return [x for x in rows if x["key"] == r["key"]]

    for r in rows:   # pass 1: segment + allowance
        hand = [x["handover"] for x in chain(r) if x["handover"]]
        r["chain_start"] = min(hand) if hand else None
        r["end"] = r["cancel"] if r["status"] == "completed" else as_of
        r["seg_start"] = r["months_basis"] = r["months_eff"] = None
        r["seg_allow"] = None
        if r["chain_start"] and r["end"]:
            # only completed rows end a segment; a completed row excludes its
            # own cancel date, an active row includes a switch on that day
            prev = [x["cancel"] for x in chain(r)
                    if x["status"] == "completed" and x["cancel"]
                    and (x["cancel"] < r["end"] if r["status"] == "completed"
                         else x["cancel"] <= r["end"])]
            seg = max([r["chain_start"]] + prev)
            end = r["end"]
            basis = datedif_months(seg, end) if end > seg else 0
            if end <= seg:
                eff = 0
            else:
                eff = basis
                if r["status"] == "completed":
                    gap = (edate(seg, basis + 1) - end).days
                    if 0 < gap <= 5:
                        eff = basis + 1
                eff += 1 if r["active"] else 0
            r.update(seg_start=seg, months_basis=basis, months_eff=eff,
                     seg_allow=eff * r["km"])

    for r in rows:   # pass 2: active rows
        r["price"] = extra_km_price(r["chain_start"], r["monthly"],
                                    r["listing"])
        r["min_left"] = (None if r["months_eff"] is None
                         else max(0, r["min_m"] - r["months_eff"]))
        r["driven"] = (None if r["current"] is None
                       else r["current"] - r["start"])
        for k in ("carry", "included", "diff", "deposit", "open_km", "days",
                  "avg_day", "avg_month", "min_end", "proj_date", "proj_km",
                  "months_end", "incl_end", "proj_diff", "cost", "proj_cost",
                  "rec_pkg"):
            r[k] = None
        if not r["active"] or r["seg_allow"] is None:
            continue
        r["carry"] = sum(x["seg_allow"] for x in chain(r)
                         if x["status"] == "completed" and x["seg_allow"])
        r["included"] = r["seg_allow"] + r["start"] + r["adj"] + r["carry"]
        if r["current"] is not None:
            r["diff"] = r["current"] - r["included"]

        plate = r["plate"].lower()
        mine = [d for d in deposits
                if d["license_plate_number"].strip().lower() == plate]
        if any(d["subscription_status"].strip() in ACTIVE for d in mine):
            r["deposit"] = sum(
                to_num(d["amount"], 0, comma=True) for d in mine
                if d["reason"].strip() == "Extra KM"
                and d["status"].strip() == "Paid"
                and d["subscription_status"].strip() in ACTIVE | {"completed"})
        else:
            r["deposit"] = 0
        if r["diff"] is not None:
            covered = (r["deposit"] / r["price"]
                       if r["deposit"] > 0 and r["price"] else 0)
            r["open_km"] = r["diff"] - covered
            if r["price"]:
                r["cost"] = max(0, r["diff"]) * r["price"]

        if r["checkpoint"] and r["chain_start"]:
            r["days"] = (r["checkpoint"] - r["chain_start"]).days
            if r["days"] > 0 and r["driven"] is not None:
                r["avg_day"] = r["driven"] / r["days"]
                r["avg_month"] = r["avg_day"] * 365.25 / 12
        r["min_end"] = edate(r["seg_start"], int(r["min_m"]))
        if r["checkpoint"]:
            r["proj_date"] = max(r["min_end"], r["checkpoint"])
        if r["avg_day"] is not None and r["proj_date"]:
            r["proj_km"] = (r["current"] + r["avg_day"]
                            * (r["proj_date"] - r["checkpoint"]).days)
        r["months_end"] = max(r["min_m"], r["months_eff"])
        r["incl_end"] = (r["start"] + r["adj"] + r["carry"]
                         + r["months_end"] * r["km"])
        if r["proj_km"] is not None:
            r["proj_diff"] = r["proj_km"] - r["incl_end"]
            if r["price"]:
                r["proj_cost"] = max(0, r["proj_diff"]) * r["price"]
        if r["avg_month"] is not None:
            avg = xround(r["avg_month"])
            r["rec_pkg"] = next((p for p in PACKAGES if p >= avg),
                                max(PACKAGES))
    return rows


def expected_web(r: dict, vat: float) -> dict:
    def i(v, div=1.0):
        return "" if v is None else xround(v / div)

    def d(v):
        return "" if v is None else serial(v)

    return {
        "vehicle": r["vehicle"], "plate": r["plate"], "status": r["status"],
        "insurance": r["insurance"], "km_package": i(r["km"]),
        "min_duration_months": i(r["min_m"]),
        "min_duration_left": i(r["min_left"]),
        "months_elapsed": i(r["months_eff"]),
        "monthly_price_chf": i(r["monthly"], vat),
        "handover_date": d(r["chain_start"]), "handover_km": i(r["start"]),
        "segment_start": d(r["seg_start"]), "min_term_end": d(r["min_end"]),
        "checkpoint_date": d(r["checkpoint"]),
        "checkpoint_type": r["checkpoint_type"],
        "current_km": i(r["current"]), "driven_km": i(r["driven"]),
        "included_km": i(r["included"]), "diff_km": i(r["diff"]),
        "extra_km_price_rp": "" if r["price"] is None
        else xround(r["price"] * 100 / vat),
        "extra_km_cost_chf": i(r["cost"], vat),
        "deposit_chf": i(r["deposit"]),           # cash: no VAT factor
        "open_extra_km": i(r["open_km"]), "avg_km_month": i(r["avg_month"]),
        "projected_km_end": i(r["proj_km"]),
        "included_km_end": i(r["incl_end"]),
        "projected_diff_end": i(r["proj_diff"]),
        "projected_cost_end_chf": i(r["proj_cost"], vat),
        "recommended_package": i(r["rec_pkg"]),
    }


# calc header token (text in the trailing parentheses) -> reference key
CALC_CHECKS = {
    "aktiv": "active", "chain_start": "chain_start", "end": "end",
    "seg_start": "seg_start", "months_basis": "months_basis",
    "months_eff": "months_eff", "seg_allowance": "seg_allow",
    "carry_over": "carry", "included_km": "included", "extra_km": "diff",
    "extra_km_price": "price", "deposit": "deposit",
    "open_extra_km": "open_km", "min_duration_left": "min_left",
    "driven_km": "driven", "days": "days", "avg_km_month": "avg_month",
    "min_term_end": "min_end", "proj_date": "proj_date",
    "projected_km_end": "proj_km", "months_at_end": "months_end",
    "included_km_end": "incl_end", "projected_diff_end": "proj_diff",
    "extra_km_cost": "cost", "projected_cost_end": "proj_cost",
    "recommended_package": "rec_pkg",
}

# Übersicht column header -> (reference key, divide by VAT)
OVERVIEW_CHECKS = {
    "Nummernschild": ("plate", False), "Fahrzeug": ("vehicle", False),
    "Aktueller Preis / Monat": ("monthly", True),
    "Mindestlaufzeit verbleibend": ("min_left", False),
    "Versicherung": ("insurance", False),
    "Extra KM Preis (CHF)": ("price", True),
    "Kilometer / Monat": ("km", False),
    "Übergabedatum": ("chain_start", False),
    "Kilometerstand Übergabe": ("start", False),
    "Datum letzter Stand": ("checkpoint", False),
    "Kilometerstand letzter Stand": ("current", False),
    "Inkludierte Kilometer": ("included", False),
    "Gefahrene Kilometer": ("driven", False),
    "Differenz": ("diff", False),
    "Mehrkilometer nach Kaution": ("open_km", False),
    "Ø km / Monat": ("avg_month", False),
    "Prognose Differenz Ende Mindestlaufzeit": ("proj_diff", False),
    "Empfohlenes Paket": ("rec_pkg", False),
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def build(out: Path, *args: str) -> Path:
    subprocess.run([sys.executable, str(BUILD), "--out", str(out),
                    "--tariffs", str(FIX_TARIFFS), *args],
                   check=True)
    return out


LOCALE_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"
 xmlns:xs="http://www.w3.org/2001/XMLSchema"
 xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Setup/L10N"><prop oor:name="ooSetupSystemLocale"
 oor:op="fuse"><value>{locale}</value></prop></item>
</oor:items>
"""


def recalc(files: list[Path], outdir: Path, locale: str | None = None) -> None:
    """Open + save with headless LibreOffice. openpyxl writes no cached
    values, so every value read back afterwards was computed by LibreOffice.
    locale (e.g. "de-DE") recalculates with that locale setting, as a proxy
    for a Google Sheet with the wrong Gebietsschema."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice (soffice) not found")
    outdir.mkdir(parents=True, exist_ok=True)
    for f in files:
        (outdir / f.name).unlink(missing_ok=True)
    prof_dir = TMP / ("lo-profile" + (f"-{locale}" if locale else ""))
    if locale:
        shutil.rmtree(prof_dir, ignore_errors=True)
        (prof_dir / "user").mkdir(parents=True)
        (prof_dir / "user" / "registrymodifications.xcu").write_text(
            LOCALE_XCU.format(locale=locale), encoding="utf-8")
    profile = prof_dir.as_uri()
    subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless",
                    "--norestore", "--convert-to", "xlsx", "--outdir",
                    str(outdir), *map(str, files)],
                   check=True, capture_output=True, timeout=300)
    for f in files:
        if not (outdir / f.name).exists():
            raise RuntimeError(f"LibreOffice did not write {f.name}")


def norm(v):
    """Normalise a cell value: blank -> "", dates -> serial, 3.0 -> 3."""
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return serial(v.date()) + (v - dt.datetime.combine(v.date(), dt.time())
                                   ).total_seconds() / 86400
    if isinstance(v, dt.date):
        return serial(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def ref_value(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, dt.date):
        return serial(v)
    return v


def same(a, b, tol=1e-6) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= tol * max(1.0, abs(b))
    return a == b


def find_errors(wb, sheets) -> list[str]:
    bad = []
    for name in sheets:
        for row in wb[name].iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and (v in ERRORS or v.startswith("Err:")):
                    bad.append(f"{name}!{c.coordinate}={v}")
    return bad


def formulas(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c.coordinate, c.value


FUNC_RE = re.compile(r"(?<![A-Za-z0-9_.!$])((?:_xlfn\.)?[A-Za-z][A-Za-z0-9.]*)\(")


def functions_in(formula: str) -> list[str]:
    no_strings = re.sub(r'"[^"]*"', '""', formula)
    return FUNC_RE.findall(no_strings)


def calc_columns(ws) -> dict[str, int]:
    cols = {}
    for c in ws[1]:
        m = re.search(r"\(([^()]*)\)\s*$", str(c.value or ""))
        if m:
            token = re.split(r"[\s:]", m.group(1).strip())[0]
            cols.setdefault(token, c.column)
    return cols


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------
SCENARIOS = {                      # name -> (as_of, vat)
    "main": (dt.date(2026, 10, 2), None),
    "vat": (dt.date(2026, 10, 2), 1.081),      # business customer
    "later": (dt.date(2027, 2, 15), None),     # other Stichtag
    "switch": (dt.date(2026, 4, 15), None),    # day of a package switch
}
VAT_COLS = {"monthly_price_chf", "extra_km_price_rp", "extra_km_cost_chf",
            "projected_cost_end_chf"}


def swap_columns(src: Path, dst: Path, a: str, b: str) -> Path:
    """Copy a fixture CSV with columns a and b swapped (header included),
    like a forked Redash query with reordered columns."""
    with open(src, newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    i, j = rows[0].index(a), rows[0].index(b)
    for row in rows:
        row[i], row[j] = row[j], row[i]
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        csv.writer(fh).writerows(rows)
    return dst


class SheetTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        TMP.mkdir(parents=True, exist_ok=True)
        src = TMP / "src"
        src.mkdir(exist_ok=True)
        cls.inputs = read_csv(FIX_INPUT)
        cls.deposits = read_csv(FIX_DEPOSITS)
        cls.built, files = {}, []
        for name, (as_of, vat) in SCENARIOS.items():
            args = ["--sample", str(FIX_INPUT), "--sample-deposits",
                    str(FIX_DEPOSITS), "--as-of", as_of.isoformat()]
            if vat:
                args += ["--vat", str(vat)]
            files.append(build(src / f"sample_{name}.xlsx", *args))
            cls.built[name] = files[-1]
        cls.template = build(src / "template.xlsx")
        files.append(cls.template)
        # broken setups: they must fail closed (blank web), not publish junk
        main_args = ["--sample-deposits", str(FIX_DEPOSITS), "--as-of",
                     SCENARIOS["main"][0].isoformat()]
        swapped = swap_columns(FIX_INPUT, src / "swapped_input.csv", "km",
                               "start_odometer")
        bad = {"swapped": build(src / "bad_swapped.xlsx", "--sample",
                                str(swapped), *main_args),
               "vat": build(src / "bad_vat.xlsx", "--sample", str(FIX_INPUT),
                            *main_args, "--vat", "1081")}
        files += bad.values()
        recalc(files, TMP / "calc")
        load = lambda d, f: openpyxl.load_workbook(d / f.name, data_only=True)
        cls.calc = {n: load(TMP / "calc", f) for n, f in cls.built.items()}
        cls.template_calc = load(TMP / "calc", cls.template)
        cls.bad = {n: load(TMP / "calc", f) for n, f in bad.items()}
        recalc([cls.built["main"]], TMP / "calc-de", locale="de-DE")
        cls.bad["locale"] = load(TMP / "calc-de", cls.built["main"])

    # --- structure -------------------------------------------------------
    def test_tabs_and_header(self):
        wb = openpyxl.load_workbook(self.template)
        self.assertEqual(wb.sheetnames, TABS)
        header = [c.value for c in wb["web"][1]]
        self.assertEqual(header, CONTRACT)
        for col, name in enumerate(CONTRACT, 1):
            for row in range(2, 12):
                fmt = wb["web"].cell(row, col).number_format
                if name in TEXT_COLS:
                    self.assertEqual(fmt, "General", name)
                else:
                    self.assertEqual(fmt, "0", f"web {name} row {row}")
        self.assertIsNone(wb["web"].cell(12, 1).value)
        cfg = wb["Config"]
        self.assertIsNone(cfg["B2"].value)
        self.assertIsNone(cfg["B3"].value)
        self.assertEqual(cfg["B4"].value, 1)
        self.assertIsNone(cfg["B5"].value)

    def test_template_has_no_data_and_imports_only_in_a1(self):
        wb = openpyxl.load_workbook(self.template)
        for name, url in (("data_input", "Config!$B$2"),
                          ("data_deposits", "Config!$B$3")):
            ws = wb[name]
            cells = [c.coordinate for row in ws.iter_rows() for c in row
                     if c.value is not None]
            self.assertEqual(cells, ["A1"], name)
            self.assertIn(f"IMPORTDATA({url})", ws["A1"].value)
            # no IFERROR: a failed fetch must stay visible in A1
            self.assertNotIn("IFERROR", ws["A1"].value.upper())
        for sheet, coord, f in formulas(wb):
            if "IMPORTDATA" in f.upper():
                self.assertIn((sheet, coord), {("data_input", "A1"),
                                               ("data_deposits", "A1")})
        for ws in wb.worksheets:   # no URLs, mails or keys in the template
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        self.assertNotRegex(c.value, r"https?://|@\w|api_key=",
                                            f"{ws.title}!{c.coordinate}")

    def test_only_allowed_functions(self):
        for path in (self.template, self.built["main"]):
            wb = openpyxl.load_workbook(path)
            n = 0
            for sheet, coord, f in formulas(wb):
                n += 1
                where = f"{sheet}!{coord}: {f[:80]}"
                for fn in functions_in(f):
                    up = fn.upper()
                    if up in ("MINIFS", "MAXIFS"):
                        self.fail(f"{fn} without _xlfn. prefix in {where}")
                    up = up.removeprefix("_XLFN.")
                    self.assertNotIn(up, FORBIDDEN_FUNCS, where)
                    self.assertIn(up, ALLOWED_FUNCS, where)
                bare = re.sub(r'"[^"]*"', '""', f)
                self.assertNotRegex(bare, r"\$?\b[A-Z]{1,3}:\$?[A-Z]{1,3}\b",
                                    f"whole-column range in {where}")
            self.assertGreater(n, 10000)

    def test_sample_typing(self):
        """--sample keeps both IMPORTDATA typings: real numbers/dates and
        ISO/decimal text, so both coercion paths are exercised."""
        ws = openpyxl.load_workbook(self.built["main"])["data_input"]
        self.assertIsInstance(ws["L2"].value, dt.datetime)        # 2026-04-15
        self.assertIsInstance(ws["J2"].value, int)
        self.assertEqual(ws["L4"].value, "2026-07-01T00:00:00")   # text
        self.assertEqual(ws["Q4"].value, "4100.0")                # text
        self.assertEqual(ws["T4"].value, "1890.0")
        dep = openpyxl.load_workbook(self.built["main"])["data_deposits"]
        self.assertEqual(dep["C3"].value, "250,5")

    # --- calculated values -----------------------------------------------
    def test_recalculated(self):
        """Values depend on Config!B5 (Stichtag), proving LibreOffice really
        computed them (openpyxl stores no cached values)."""
        for name, (as_of, _vat) in SCENARIOS.items():
            self.assertEqual(norm(self.calc[name]["Config"]["B6"].value),
                             serial(as_of), name)
        today = norm(self.template_calc["Config"]["B6"].value)
        self.assertEqual(today, serial(dt.date.today()))

    def test_no_errors(self):
        for name, wb in self.calc.items():
            self.assertEqual(find_errors(wb, ["calc", "web", "Übersicht",
                                              "Config"]), [], name)
        self.assertEqual(find_errors(self.template_calc,
                                     ["calc", "web", "Übersicht"]), [])

    def test_template_blank(self):
        wb = self.template_calc
        for row in wb["web"].iter_rows(min_row=2, max_row=11, values_only=True):
            self.assertEqual([norm(v) for v in row], [""] * len(CONTRACT))
        for row in wb["Übersicht"].iter_rows(min_row=4, max_row=53,
                                             values_only=True):
            self.assertEqual({norm(v) for v in row}, {""})
        for row in wb["calc"].iter_rows(min_row=2, max_row=201,
                                        values_only=True):
            self.assertEqual({norm(v) for v in row}, {""})
        self.assertEqual(wb["Config"]["B9"].value,
                         "Trag zuerst die Redash-URL in B2 ein.")
        self.assertEqual(wb["Config"]["B10"].value,
                         "Keine Kautions-URL (optional).")
        self.assertEqual(norm(wb["Config"]["D9"].value), 0)

    def test_web_matches_reference(self):
        for name, (as_of, vat) in SCENARIOS.items():
            ref = [r for r in reference(self.inputs, self.deposits, as_of)
                   if r["active"]]
            self.assertGreaterEqual(len(ref), 5)
            ws = self.calc[name]["web"]
            self.assertEqual([c.value for c in ws[1]], CONTRACT)
            for i in range(10):
                actual = [norm(c.value) for c in ws[i + 2]][:len(CONTRACT)]
                if i < len(ref):
                    exp = expected_web(ref[i], vat or 1.0)
                    expected = [exp[k] for k in CONTRACT]
                else:
                    expected = [""] * len(CONTRACT)
                for k, a, e in zip(CONTRACT, actual, expected):
                    self.assertEqual(a, e, f"{name} web row {i + 2} {k}")
                    if a != "" and k not in TEXT_COLS:
                        self.assertIsInstance(a, int, f"{name} {k}")

    def test_calc_matches_reference(self):
        for name, (as_of, _vat) in SCENARIOS.items():
            ws = self.calc[name]["calc"]
            cols = calc_columns(ws)
            for r in reference(self.inputs, self.deposits, as_of):
                for token, key in CALC_CHECKS.items():
                    a = norm(ws.cell(r["sheet_row"], cols[token]).value)
                    e = ref_value(r[key])
                    self.assertTrue(same(a, e), f"{name} calc row "
                                    f"{r['sheet_row']} {token}: {a!r} != {e!r}")
            ranks = [norm(ws.cell(row, cols["aktiv_rang"]).value)
                     for row in range(2, 202)]
            n_active = sum(1 for r in self.inputs
                           if r["subscription_status"].strip() in ACTIVE)
            self.assertEqual([x for x in ranks if x != ""],
                             list(range(1, n_active + 1)))

    def test_overview_matches_reference(self):
        for name, (as_of, vat) in SCENARIOS.items():
            ws = self.calc[name]["Übersicht"]
            headers = [c.value for c in ws[3]]
            self.assertEqual(headers, list(OVERVIEW_CHECKS))
            ref = [r for r in reference(self.inputs, self.deposits, as_of)
                   if r["active"]]
            for i in range(50):
                row = [norm(c.value) for c in ws[i + 4]][:len(headers)]
                for h, a in zip(headers, row):
                    if i >= len(ref):
                        self.assertEqual(a, "", f"{name} Übersicht {i + 4} {h}")
                        continue
                    key, div = OVERVIEW_CHECKS[h]
                    e = ref_value(ref[i][key])
                    if div and e != "":
                        e = e / (vat or 1.0)
                    self.assertTrue(same(a, e), f"{name} Übersicht row {i + 4} "
                                    f"{h}: {a!r} != {e!r}")

    def test_fixture_coverage(self):
        """The fixtures cover all required business cases."""
        ref = reference(self.inputs, self.deposits, SCENARIOS["main"][0])
        by_status = {r["status"] for r in ref}
        self.assertTrue({"on_lease", "notice_given", "completed"} <= by_status)
        self.assertTrue(any(r["carry"] for r in ref))                  # (a)
        self.assertTrue(any(r["status"] == "completed" and r["months_eff"]
                            == r["months_basis"] + 1 for r in ref))     # (d)
        self.assertTrue(any(r["chain_start"] < CUTOFF and r["active"]
                            for r in ref))                             # (e)
        self.assertTrue(any(r["chain_start"] >= CUTOFF and r["listing"]
                            and r["listing"] > 100000 for r in ref))
        self.assertTrue(any(r["adj"] for r in ref))                    # (f)
        self.assertTrue(any(r["active"] and r["current"] is None
                            for r in ref))                             # (g)
        self.assertTrue(any(r["deposit"] for r in ref))                # (h)
        self.assertTrue(any("T" in x["handover_date"] for x in self.inputs))
        self.assertTrue(any(r["active"] and r["cancel"]                # (i)
                            and r["cancel"] < SCENARIOS["main"][0]
                            for r in ref))
        switch = SCENARIOS["switch"][0]                                # (j)
        self.assertTrue(any(r["status"] == "completed" and r["cancel"] == switch
                            for r in ref))

    # --- regression values (hard-coded, independent of reference()) -------
    def web_row(self, wb, plate: str) -> dict:
        for row in wb["web"].iter_rows(min_row=2, max_row=11):
            vals = {k: norm(c.value) for k, c in zip(CONTRACT, row)}
            if vals["plate"] == plate:
                return vals
        self.fail(f"{plate} not in web")

    def test_active_row_with_past_cancellation(self):
        """notice_given row whose cancellation_date already passed: its own
        date must not start a new km segment."""
        r = self.web_row(self.calc["main"], "XX000008")
        self.assertEqual(
            {k: r[k] for k in ("segment_start", "months_elapsed",
                               "min_duration_left", "min_term_end",
                               "included_km", "diff_km", "extra_km_cost_chf")},
            {"segment_start": serial(dt.date(2025, 1, 10)),
             "months_elapsed": 21, "min_duration_left": 0,
             "min_term_end": serial(dt.date(2026, 1, 10)),
             # 3900 km x fake old tariff for a monthly price <= 900 (0.40)
             "included_km": 21100, "diff_km": 3900, "extra_km_cost_chf": 1560})

    def test_switch_day_no_double_count(self):
        """On the day of a package switch (Stichtag = cancellation_date of
        the completed row) the new segment starts that day."""
        r = self.web_row(self.calc["switch"], "XX000001")
        self.assertEqual(
            {k: r[k] for k in ("segment_start", "months_elapsed",
                               "included_km", "diff_km")},
            {"segment_start": serial(dt.date(2026, 4, 15)),
             "months_elapsed": 0, "included_km": 20420, "diff_km": 580})

    # --- broken setups fail closed ----------------------------------------
    def assert_web_blank(self, wb, name):
        for row in wb["web"].iter_rows(min_row=2, max_row=11, values_only=True):
            self.assertEqual([norm(v) for v in row], [""] * len(CONTRACT), name)
        self.assertEqual([c.value for c in wb["web"][1]], CONTRACT, name)

    def test_wrong_locale_blocks_web(self):
        wb = self.bad["locale"]
        self.assert_web_blank(wb, "locale")
        for ref in ("B9", "B10"):
            self.assertTrue(str(wb["Config"][ref].value).startswith(
                "Achtung: Gebietsschema falsch"), ref)
        self.assertEqual(norm(wb["Config"]["D9"].value), 0)
        self.assertEqual(wb["Übersicht"]["N1"].value, wb["Config"]["B9"].value)

    def test_wrong_columns_block_web(self):
        wb = self.bad["swapped"]
        self.assert_web_blank(wb, "swapped")
        self.assertTrue(str(wb["Config"]["B9"].value).startswith(
            "Achtung: Spalten in data_input passen nicht"))
        self.assertEqual(norm(wb["Config"]["D9"].value), 0)
        self.assertEqual(find_errors(wb, ["calc", "web", "Übersicht"]), [])

    def test_invalid_vat_blanks_prices(self):
        wb = self.bad["vat"]
        self.assertEqual(wb["Config"]["B7"].value, "#N/A")
        self.assertTrue(str(wb["Übersicht"]["K1"].value).startswith(
            "MwSt-Faktor ungültig"))
        self.assertEqual(find_errors(wb, ["calc", "web", "Übersicht"]), [])
        ref = [r for r in reference(self.inputs, self.deposits,
                                    SCENARIOS["main"][0]) if r["active"]]
        for i, r in enumerate(ref):
            exp = expected_web(r, 1.0)
            row = [norm(c.value) for c in wb["web"][i + 2]]
            for k, a in zip(CONTRACT, row):
                self.assertEqual(a, "" if k in VAT_COLS else exp[k],
                                 f"bad vat web row {i + 2} {k}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
