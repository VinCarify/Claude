#!/usr/bin/env python3
"""Generator for the CARIFY KM-Tracker Google-Sheets template.

The workbook is written as .xlsx and later uploaded to Google Drive, where it
is converted to a Google Sheet. Formulas therefore stick to classic functions
that both Google Sheets and LibreOffice (used by the tests) understand.

    python3 build_sheet.py --out km-tracker-template.xlsx [--tariffs t.json]
    python3 build_sheet.py --out sample.xlsx --sample fixtures/sample_input.csv \
        --sample-deposits fixtures/sample_deposits.csv --as-of 2026-10-02

Without --sample the template contains no data, only the IMPORTDATA formulas
in data_input!A1 and data_deposits!A1. --sample writes fixture rows as static
values instead (testing only).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------
FIRST, LAST = 2, 201          # data rows of data_input / calc
DEP_LAST = 2000               # last data row of data_deposits
WEB_ROWS = 10                 # rows 2..11 in "web"
OVERVIEW_ROWS = 50            # active subscriptions listed in "Übersicht"

INPUT_COLUMNS = [
    "vehicle", "subscription_status", "first_name", "last_name",
    "minimum_duration_months", "insurance_package_type", "cancellation_date",
    "license_plate_number", "vehicle_id", "start_odometer", "s_listing",
    "handover_date", "km", "vin_number", "total_km_adjustment",
    "checkpoint_date", "current_odometer", "checkpoint_type",
    "deposit_amount_total", "monthly_price",
]
DEPOSIT_COLUMNS = [
    "deposit_id", "time_dim", "amount", "payment_due_date", "refunded_amount",
    "refund_date", "subscription_id", "subscription_status", "return_date",
    "customer_name", "company_name", "stripe_invoice_id", "reference",
    "comments", "status", "reason", "license_plate_number",
]
CONTRACT = [
    "vehicle", "plate", "status", "insurance", "km_package",
    "min_duration_months", "min_duration_left", "months_elapsed",
    "monthly_price_chf", "handover_date", "handover_km", "segment_start",
    "min_term_end", "checkpoint_date", "checkpoint_type", "current_km",
    "driven_km", "included_km", "diff_km", "extra_km_price_rp",
    "extra_km_cost_chf", "deposit_chf", "open_extra_km", "avg_km_month",
    "projected_km_end", "included_km_end", "projected_diff_end",
    "projected_cost_end_chf", "recommended_package",
]

# Tariffs are company data and live outside the public repo: pass --tariffs
# or put tariffs.local.json (gitignored) next to this script. Format:
# {"cutoff": "YYYY-MM-DD",
#  "old": [[upper_monthly_price, chf_per_km], ..., [null, null]],
#  "new": [[upper_list_price, chf_per_km], ..., [null, chf_per_km]],
#  "packages": [500, 750, ...]}            (packages optional)
# Upper bounds are inclusive; a null bound is the "darüber" row.
LOCAL_TARIFFS = Path(__file__).with_name("tariffs.local.json")
DEFAULT_PACKAGES = [500, 750, 1000, 1750, 2000, 2500, 3000, 3500, 4000, 5000]
TARIFF_ROWS = (6, 5)          # fixed table sizes (old, new) – formulas rely on it


def load_tariffs(path: str | None) -> dict:
    src = Path(path) if path else LOCAL_TARIFFS
    if not src.exists():
        raise SystemExit(f"Tarif-Datei fehlt: {src} (siehe Kommentar in "
                         "build_sheet.py)")
    t = json.loads(src.read_text(encoding="utf-8"))
    old = [tuple(r) for r in t["old"]]
    new = [tuple(r) for r in t["new"]]
    if (len(old), len(new)) != TARIFF_ROWS or old[-1][0] is not None \
            or new[-1][0] is not None:
        raise SystemExit(f"Tarif-Datei {src}: erwartet {TARIFF_ROWS[0]} alte "
                         f"und {TARIFF_ROWS[1]} neue Zeilen, letzte mit null")
    return {"cutoff": dt.date.fromisoformat(t["cutoff"]), "old": old,
            "new": new, "packages": t.get("packages", DEFAULT_PACKAGES)}


PACKAGE_SLOTS = 15            # room for a few more packages

# Config layout (row numbers are referenced by formulas)
CFG_ASOF = "Config!$B$6"
CFG_VAT = "Config!$B$7"
CFG_CUTOFF = "Config!$B$8"
CFG_KM_OK = "Config!$D$9"     # 1 = data_input usable (locale + header)
CFG_DEP_OK = "Config!$D$10"   # 1 = data_deposits usable (locale + header)
# "1234.5" is no date in any locale; only a "." decimal locale reads it right
LOCALE_OK = 'IFERROR(VALUE("1234.5"),0)=1234.5'
LOCALE_MSG = ("Achtung: Gebietsschema falsch. Datei → Einstellungen → "
              "Gebietsschema «Schweiz». Bis dahin bleibt «web» leer.")
OLD_ROW = 14                  # first row of the old tariff table
NEW_ROW = OLD_ROW + TARIFF_ROWS[0] + 3
PKG_ROW = NEW_ROW + TARIFF_ROWS[1] + 2
OLD_UP = f"Config!$A${OLD_ROW}:$A${OLD_ROW + TARIFF_ROWS[0] - 1}"
OLD_PRICE = f"Config!$B${OLD_ROW}:$B${OLD_ROW + TARIFF_ROWS[0] - 1}"
NEW_UP = f"Config!$A${NEW_ROW}:$A${NEW_ROW + TARIFF_ROWS[1] - 1}"
NEW_PRICE = f"Config!$B${NEW_ROW}:$B${NEW_ROW + TARIFF_ROWS[1] - 1}"
PKG_RNG = f"Config!$A${PKG_ROW}:$A${PKG_ROW + PACKAGE_SLOTS - 1}"

# Styling
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16)
HEAD_FILL = PatternFill("solid", fgColor="E8F0FE")
GROUP_FILL = PatternFill("solid", fgColor="C9DAF8")
NOTE_FONT = Font(italic=True, color="666666")
THIN = Side(style="thin", color="B7C4D6")
HEAD_BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER_WRAP = Alignment(wrap_text=True, horizontal="center", vertical="center")

FMT_DATE = "dd.mm.yyyy"
FMT_KM = "#,##0"
FMT_CHF = "#,##0.00"
FMT_INT = "0"


# --------------------------------------------------------------------------
# Formula helpers
# --------------------------------------------------------------------------
def num(ref: str, default: str) -> str:
    """Number or numeric text ("1399.0") -> number; blank -> default."""
    return (f'IF(ISNUMBER({ref}),{ref},IF(LEN(TRIM({ref}))=0,{default},'
            f'IFERROR(VALUE(TRIM({ref})),{default})))')


def date(ref: str) -> str:
    """Date serial or ISO text ("2025-08-01", "2025-08-01T00:00:00") -> whole
    day serial; blank -> "". ROUND(x-0.5,0) drops a time-of-day part."""
    t = f"TRIM({ref})"
    return (f'IF(ISNUMBER({ref}),ROUND({ref}-0.5,0),IF(LEN({t})<10,"",'
            f'IFERROR(DATE(VALUE(LEFT({t},4)),VALUE(MID({t},6,2)),'
            f'VALUE(MID({t},9,2))),"")))')


def di(col: str) -> str:
    return f"data_input!${col}{{r}}"


def dd(col: str) -> str:
    return f"data_deposits!${col}{{r}}"


# calc columns: (key, header, formula expression, number format)
# {key} = same-row cell, {key_rng} = rows 2..201 (abs.), {key_first} = $X$2,
# {key_drng} = rows 2..2000 (deposit helper block).
CALC = [
    # Import kaputt (Gebietsschema, Spalten) = alles leer, nichts Falsches raus
    ("has_row", "Zeile belegt + Import ok (has_row)",
     f'IF(AND({CFG_KM_OK}=1,LEN(TRIM(data_input!$B{{r}}))>0),1,"")', None),
    ("active", "Aktiv (aktiv: on_lease / notice_given)",
     'IF(OR({status}="on_lease",{status}="notice_given"),1,0)', None),
    ("rank", "Aktiv-Rang (aktiv_rang)",
     'IF({active}=1,COUNTIF({active_first}:{active},1),"")', None),
    ("plate", "Nummernschild (license_plate_number)",
     f'TRIM({di("H")}&"")', None),
    ("vehicle", "Fahrzeug (vehicle)", f'TRIM({di("A")}&"")', None),
    ("status", "Status (subscription_status)", f'TRIM({di("B")}&"")', None),
    ("last_name", "Nachname (last_name)", f'TRIM({di("D")}&"")', None),
    ("chain_key", "Ketten-Schlüssel (chain_key = last_name|plate)",
     '{last_name}&"|"&{plate}', None),
    ("insurance", "Versicherung (insurance_package_type)",
     f'TRIM({di("F")}&"")', None),
    ("checkpoint_type", "Typ letzter Stand (checkpoint_type)",
     f'TRIM({di("R")}&"")', None),
    ("km_package", "KM-Paket / Monat (km)", num(di("M"), "0"), FMT_KM),
    ("min_months", "Mindestlaufzeit Monate (minimum_duration_months)",
     num(di("E"), "0"), FMT_INT),
    ("start_km", "Kilometerstand Übergabe (start_odometer)",
     num(di("J"), "0"), FMT_KM),
    ("adjustment", "KM-Korrektur (total_km_adjustment)",
     num(di("O"), "0"), FMT_KM),
    ("current_km", "Kilometerstand letzter Stand (current_odometer)",
     num(di("Q"), '""'), FMT_KM),
    ("listing", "Listenpreis CHF (s_listing)", num(di("K"), '""'), FMT_CHF),
    ("monthly", "Monatspreis CHF inkl. MwSt (monthly_price)",
     num(di("T"), '""'), FMT_CHF),
    ("handover", "Übergabe dieses Abos (handover_date)", date(di("L")),
     FMT_DATE),
    ("cancel", "Kündigungsdatum (cancellation_date)", date(di("G")), FMT_DATE),
    ("checkpoint", "Datum letzter Stand (checkpoint_date)", date(di("P")),
     FMT_DATE),
    ("chain_start", "Abo-Start Kette (chain_start)",
     'IF(COUNTIFS({chain_key_rng},{chain_key},{handover_rng},">0")=0,"",'
     '_xlfn.MINIFS({handover_rng},{chain_key_rng},{chain_key}))', FMT_DATE),
    ("end", "Enddatum Berechnung (end = Kündigung oder Stichtag)",
     f'IF({{status}}="completed",{{cancel}},{CFG_ASOF})', FMT_DATE),
    # Nur abgeschlossene Abos der Kette beenden ein Segment. Abgeschlossene
    # Zeile: Kündigung < end (eigene ausschliessen); aktive: <= end (am Tag
    # der Verlängerung startet das neue Segment schon).
    ("seg_start", "Start KM-Paket-Segment (seg_start)",
     'IF(OR({chain_start}="",{end}=""),"",MAX({chain_start},'
     '_xlfn.MAXIFS({cancel_rng},{chain_key_rng},{chain_key},'
     '{status_rng},"completed",'
     '{cancel_rng},IF({status}="completed","<","<=")&{end})))', FMT_DATE),
    ("months_basis", "Volle Monate im Segment (months_basis)",
     'IF({seg_start}="","",IF({end}>{seg_start},'
     'DATEDIF({seg_start},{end},"M"),0))', FMT_INT),
    # 5 Tage Kulanz bei abgeschlossenen Abos, laufender Monat zählt bei aktiven
    ("months_eff", "Monate verrechnet (months_eff)",
     'IF({months_basis}="","",IF({end}<={seg_start},0,'
     'IF({status}="completed",IF(AND(EDATE({seg_start},{months_basis}+1)-{end}>0,'
     'EDATE({seg_start},{months_basis}+1)-{end}<=5),{months_basis}+1,'
     '{months_basis}),{months_basis})+IF({active}=1,1,0)))', FMT_INT),
    ("seg_allow", "KM-Guthaben Segment (seg_allowance)",
     'IF({months_eff}="","",{months_eff}*{km_package})', FMT_KM),
    ("carry", "Übertrag abgeschlossene Abos (carry_over)",
     'IF({active}<>1,"",SUMIFS({seg_allow_rng},{chain_key_rng},{chain_key},'
     '{status_rng},"completed"))', FMT_KM),
    ("included", "Inkludierte KM als Kilometerstand (included_km)",
     'IF(OR({active}<>1,{seg_allow}=""),"",'
     '{seg_allow}+{start_km}+{adjustment}+{carry})', FMT_KM),
    ("diff", "Differenz / Mehrkilometer (extra_km = diff_km)",
     'IF(OR({included}="",{current_km}=""),"",{current_km}-{included})',
     FMT_KM),
    ("tariff", "Tarif (alt = nach Monatspreis, neu = nach Listenpreis)",
     f'IF({{chain_start}}="","",IF({{chain_start}}<{CFG_CUTOFF},"alt","neu"))',
     None),
    ("price_lookup", "Tarif-Lookup CHF/km (0 = kein Tarif)",
     f'IF({{tariff}}="alt",IF({{monthly}}="","",N(INDEX({OLD_PRICE},'
     f'COUNTIF({OLD_UP},"<"&{{monthly}})+1))),IF({{tariff}}="neu",'
     f'IF({{listing}}="","",N(INDEX({NEW_PRICE},'
     f'COUNTIF({NEW_UP},"<"&{{listing}})+1))),""))', "0.00"),
    ("price", "Extra-KM-Preis CHF/km inkl. MwSt (extra_km_price)",
     'IF(N({price_lookup})>0,{price_lookup},"")', "0.00"),
    # Kautions-URL gesetzt, aber Import kaputt = leer statt still 0
    ("deposit", "Kaution Extra KM CHF (deposit)",
     f'IF({{active}}<>1,"",IF(AND(LEN(Config!$B$3)>0,{CFG_DEP_OK}<>1),"",'
     'IF(COUNTIFS({dep_plate_drng},{plate},'
     '{dep_active_drng},1)>0,SUMIFS({dep_amount_drng},{dep_plate_drng},'
     '{plate},{dep_counts_drng},1),0)))', FMT_CHF),
    ("open_km", "Mehrkilometer nach Kaution (open_extra_km)",
     'IF(OR({diff}="",{deposit}=""),"",{diff}-IF(AND(N({deposit})>0,'
     'N({price})>0),N({deposit})/N({price}),0))', FMT_KM),
    ("min_left", "Mindestlaufzeit verbleibend Monate (min_duration_left)",
     'IF({months_eff}="","",MAX(0,{min_months}-{months_eff}))', FMT_INT),
    ("driven", "Gefahrene KM (driven_km)",
     'IF({current_km}="","",{current_km}-{start_km})', FMT_KM),
    ("days", "Tage Abo-Start bis letzter Stand (days)",
     'IF(OR({active}<>1,{checkpoint}="",{chain_start}=""),"",'
     '{checkpoint}-{chain_start})', FMT_INT),
    ("avg_day", "Ø km / Tag (avg_per_day)",
     'IF(OR({days}="",{driven}=""),"",IF({days}>0,{driven}/{days},""))',
     "0.00"),
    ("avg_month", "Ø km / Monat (avg_km_month)",
     'IF({avg_day}="","",{avg_day}*365.25/12)', FMT_KM),
    ("min_end", "Ende Mindestlaufzeit (min_term_end)",
     'IF(OR({active}<>1,{seg_start}=""),"",EDATE({seg_start},{min_months}))',
     FMT_DATE),
    ("proj_date", "Prognose-Datum (proj_date)",
     'IF(OR({min_end}="",{checkpoint}=""),"",MAX({min_end},{checkpoint}))',
     FMT_DATE),
    ("proj_km", "Prognose Kilometerstand Ende (projected_km_end)",
     'IF(OR({avg_day}="",{proj_date}=""),"",'
     '{current_km}+{avg_day}*({proj_date}-{checkpoint}))', FMT_KM),
    ("months_end", "Monate bis Ende (months_at_end)",
     'IF(OR({active}<>1,{months_eff}=""),"",MAX({min_months},{months_eff}))',
     FMT_INT),
    ("incl_end", "Inkludierte KM Ende (included_km_end)",
     'IF({months_end}="","",'
     '{start_km}+{adjustment}+{carry}+{months_end}*{km_package})', FMT_KM),
    ("proj_diff", "Prognose Differenz Ende (projected_diff_end)",
     'IF(OR({proj_km}="",{incl_end}=""),"",{proj_km}-{incl_end})', FMT_KM),
    ("cost", "Kosten Mehrkilometer CHF inkl. MwSt (extra_km_cost)",
     'IF(OR({diff}="",{price}=""),"",MAX(0,{diff})*{price})', FMT_CHF),
    ("proj_cost", "Prognose Kosten Ende CHF inkl. MwSt (projected_cost_end)",
     'IF(OR({proj_diff}="",{price}=""),"",MAX(0,{proj_diff})*{price})',
     FMT_CHF),
    ("rec_pkg", "Empfohlenes Paket km / Monat (recommended_package)",
     f'IF({{avg_month}}="","",IF(COUNTIF({PKG_RNG},">="&ROUND({{avg_month}},0))>0,'
     f'_xlfn.MINIFS({PKG_RNG},{PKG_RNG},">="&ROUND({{avg_month}},0)),'
     f'MAX({PKG_RNG})))', FMT_KM),
]

# Deposit helper block in calc (one row per data_deposits row 2..2000),
# placed right of the main block after one empty separator column.
DEPOSIT_HELPER = [
    ("dep_plate", "Kaution: Nummernschild (license_plate_number)",
     f'IF(OR({CFG_DEP_OK}<>1,LEN(TRIM({dd("Q")}))=0),"",'
     f'TRIM({dd("Q")}&""))'),
    ("dep_amount", "Kaution: Betrag CHF (amount)",
     f'IF({{dep_plate}}="","",IFERROR(IF(ISNUMBER({dd("C")}),{dd("C")},'
     f'VALUE(SUBSTITUTE(TRIM({dd("C")}),",","."))),0))'),
    ("dep_substatus", "Kaution: Abo-Status (subscription_status)",
     f'IF({{dep_plate}}="","",TRIM({dd("H")}&""))'),
    ("dep_active", "Kaution: Abo aktiv (on_lease / notice_given)",
     'IF({dep_plate}="","",IF(OR({dep_substatus}="on_lease",'
     '{dep_substatus}="notice_given"),1,0))'),
    ("dep_counts", "Kaution: zählt (Extra KM + Paid)",
     f'IF({{dep_plate}}="","",IF(AND(TRIM({dd("P")}&"")="Extra KM",'
     f'TRIM({dd("O")}&"")="Paid",OR({{dep_substatus}}="on_lease",'
     f'{{dep_substatus}}="notice_given",{{dep_substatus}}="completed")),1,0))'),
]

COL = {key: get_column_letter(i + 1) for i, (key, *_r) in enumerate(CALC)}
_dep_start = len(CALC) + 2
COL.update({key: get_column_letter(_dep_start + i)
            for i, (key, *_r) in enumerate(DEPOSIT_HELPER)})


class _Refs(dict):
    """format_map helper resolving {key}, {key_rng}, {key_drng}, {key_first}."""

    def __init__(self, row: int):
        super().__init__(r=row)
        self.row = row

    def __missing__(self, name: str) -> str:
        for suffix, (a, b) in (("_drng", (FIRST, DEP_LAST)),
                               ("_rng", (FIRST, LAST))):
            if name.endswith(suffix):
                c = COL[name[: -len(suffix)]]
                return f"${c}${a}:${c}${b}"
        if name.endswith("_first"):
            return f"${COL[name[:-6]]}${FIRST}"
        return f"{COL[name]}{self.row}"


def calc_formula(key: str, expr: str, row: int) -> str:
    body = expr.format_map(_Refs(row))
    if key == "has_row":
        return "=" + body
    return f'=IF({COL["has_row"]}{row}="","",IFERROR({body},""))'


def calc_ref(key: str) -> str:
    c = COL[key]
    return f"calc!${c}${FIRST}:${c}${LAST}"


def pick(key: str, rank: int) -> str:
    """INDEX/MATCH of a calc column for the given active rank."""
    return (f"INDEX({calc_ref(key)},MATCH({rank},{calc_ref('rank')},0))")


# web: contract key -> (calc key, kind); kinds: text | int | vat | rp
WEB = {
    "vehicle": ("vehicle", "text"), "plate": ("plate", "text"),
    "status": ("status", "text"), "insurance": ("insurance", "text"),
    "km_package": ("km_package", "int"),
    "min_duration_months": ("min_months", "int"),
    "min_duration_left": ("min_left", "int"),
    "months_elapsed": ("months_eff", "int"),
    "monthly_price_chf": ("monthly", "vat"),
    "handover_date": ("chain_start", "int"),
    "handover_km": ("start_km", "int"),
    "segment_start": ("seg_start", "int"),
    "min_term_end": ("min_end", "int"),
    "checkpoint_date": ("checkpoint", "int"),
    "checkpoint_type": ("checkpoint_type", "text"),
    "current_km": ("current_km", "int"), "driven_km": ("driven", "int"),
    "included_km": ("included", "int"), "diff_km": ("diff", "int"),
    "extra_km_price_rp": ("price", "rp"),
    "extra_km_cost_chf": ("cost", "vat"),
    "deposit_chf": ("deposit", "int"),       # cash held: no VAT factor
    "open_extra_km": ("open_km", "int"),
    "avg_km_month": ("avg_month", "int"),
    "projected_km_end": ("proj_km", "int"),
    "included_km_end": ("incl_end", "int"),
    "projected_diff_end": ("proj_diff", "int"),
    "projected_cost_end_chf": ("proj_cost", "vat"),
    "recommended_package": ("rec_pkg", "int"),
}
assert list(WEB) == CONTRACT

# Übersicht: (group, header, calc key, kind, number format, width)
OVERVIEW = [
    ("Fahrzeug", "Nummernschild", "plate", "text", None, 15),
    ("Fahrzeug", "Fahrzeug", "vehicle", "text", None, 30),
    ("Abo", "Aktueller Preis / Monat", "monthly", "vat", FMT_CHF, 13),
    ("Abo", "Mindestlaufzeit verbleibend", "min_left", "raw", '0 "Mt."', 14),
    ("Abo", "Versicherung", "insurance", "text", None, 14),
    ("Abo", "Extra KM Preis (CHF)", "price", "vat", FMT_CHF, 11),
    ("Abo", "Kilometer / Monat", "km_package", "raw", FMT_KM, 11),
    ("Übergabe", "Übergabedatum", "chain_start", "raw", FMT_DATE, 12),
    ("Übergabe", "Kilometerstand Übergabe", "start_km", "raw", FMT_KM, 14),
    ("Letzter Stand", "Datum letzter Stand", "checkpoint", "raw", FMT_DATE, 12),
    ("Letzter Stand", "Kilometerstand letzter Stand", "current_km", "raw",
     FMT_KM, 14),
    ("Kilometer-Bilanz", "Inkludierte Kilometer", "included", "raw", FMT_KM, 12),
    ("Kilometer-Bilanz", "Gefahrene Kilometer", "driven", "raw", FMT_KM, 12),
    ("Kilometer-Bilanz", "Differenz", "diff", "raw", FMT_KM, 11),
    ("Kilometer-Bilanz", "Mehrkilometer nach Kaution", "open_km", "raw",
     FMT_KM, 15),
    ("Prognose", "Ø km / Monat", "avg_month", "raw", FMT_KM, 11),
    ("Prognose", "Prognose Differenz Ende Mindestlaufzeit", "proj_diff", "raw",
     FMT_KM, 15),
    ("Prognose", "Empfohlenes Paket", "rec_pkg", "raw", FMT_KM, 12),
]


# --------------------------------------------------------------------------
# Sheet builders
# --------------------------------------------------------------------------
def build_anleitung(ws) -> None:
    ws["A1"] = "CARIFY KM-Tracker – so gehts"
    ws["A1"].font = TITLE
    ws["A2"] = ("Hey! Mit diesem Sheet trackst du die Kilometer von Abos und "
                "fütterst das Kunden-Dashboard. Einmal einrichten, dann läuft "
                "alles von allein.")
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:B2")

    rows = [
        ("Setup in 5 Schritten", None),
        ("1", "Redash: Forke die bestehende Fleet-Query. Setz den Filter auf "
              "die gewünschte(n) Abo-Nummer(n), z. B. subscription_number IN "
              "(...). Stell den Refresh-Schedule auf täglich. Dann kopierst du "
              "die CSV-Results-URL (die mit api_key)."),
        ("2", "Füg die URL in Config!B2 ein. Optional: die CSV-URL der "
              "Kautions-Query in Config!B3."),
        ("3", "Check den MwSt-Faktor in Config!B4: 1 = Privatkunde (Preise "
              "inkl. MwSt), 1.081 = Firmenkunde (Preise exkl. MwSt)."),
        ("4", "Datei → Freigeben → Im Web veröffentlichen → nur Tab «web» → "
              "CSV → Link kopieren."),
        ("5", "Hinterleg den Link als SHEET_CSV_URL in Netlify (Environment "
              "variables). Fertig!"),
        (None, None),
        ("Wichtig", None),
        ("!", "Datei → Einstellungen: Gebietsschema «Schweiz» und Zeitzone "
              "«Zürich» setzen. Sonst liest IMPORTDATA Zahlen und Daten "
              "falsch. Im Tab «Berechnung» die Neuberechnung auf «Bei "
              "Änderung und jede Stunde» stellen. So stimmt «heute» immer."),
        ("!", "Tab «web» nicht umbauen. Spaltennamen und Reihenfolge sind der "
              "Vertrag mit der Website."),
        ("!", "Der Web-Link ist öffentlich. Darum stehen im Tab «web» keine "
              "Namen. Lass das so und teil das Sheet selbst nur intern."),
        ("!", "In data_input und data_deposits schreibt nur der Import. Nichts "
              "reintippen, sonst bricht der Import ab (#REF!)."),
        ("!", "Config!B5 (Stichtag) leer lassen = heute. Nur für Tests oder "
              "Rückblicke ausfüllen."),
        ("!", "Config!B9 und B10 zeigen dir, ob Import und Gebietsschema "
              "passen. Wenn nicht, bleibt «web» leer. So landen keine falschen "
              "Zahlen auf der Website."),
        (None, None),
        ("So rechnet das Sheet", None),
        ("Kette", "Gleicher Nachname + gleiches Nummernschild = eine Kette. "
                  "Verlängerung oder Paketwechsel starten ein neues Abo, das "
                  "alte wird «completed»."),
        ("Inkludiert", "Kilometerstand bei Übergabe + Korrektur + Guthaben aus "
                       "abgeschlossenen Abos + Monate × Paket. Bei aktiven Abos "
                       "zählt der laufende Monat voll. Bei abgeschlossenen gibts "
                       "5 Tage Kulanz."),
        ("Differenz", "Aktueller Kilometerstand minus inkludierte km. Positiv = "
                      "Mehrkilometer, negativ = Puffer."),
        ("Extra-KM-Preis", "Abo-Start vor dem Tarifwechsel (Config!B8): nach "
                           "Monatspreis. Danach: nach Listenpreis. Die Tabellen "
                           "findest du in Config."),
        ("Kaution", "Bezahlte Extra-KM-Kautionen werden in km umgerechnet und "
                    "von den Mehrkilometern abgezogen."),
        ("Prognose", "Ø km pro Tag seit Abo-Start, hochgerechnet bis zum Ende "
                     "der Mindestlaufzeit."),
        (None, None),
        ("Die Tabs", None),
        ("Übersicht", "Alle aktiven Abos auf einen Blick. Fürs Team."),
        ("web", "Daten für die Website (max. 10 aktive Abos)."),
        ("calc", "Alle Zwischenschritte. Hier rechnet das Sheet."),
        ("data_input", "Rohdaten KM aus Redash (via IMPORTDATA)."),
        ("data_deposits", "Rohdaten Kautionen aus Redash (optional)."),
        ("Config", "Einstellungen, Extra-KM-Tarife und KM-Pakete."),
    ]
    r = 4
    for a, b in rows:
        if a is not None:
            ws.cell(r, 1, a)
        if b is not None:
            ws.cell(r, 2, b).alignment = WRAP
        if a is not None and b is None:
            ws.cell(r, 1).font = Font(bold=True, size=12)
        else:
            ws.cell(r, 1).font = BOLD
            ws.cell(r, 1).alignment = Alignment(vertical="top")
        r += 1
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 95
    ws.row_dimensions[2].height = 32


def build_config(ws, vat: float | None, as_of: dt.date | None,
                 tariffs: dict) -> None:
    ws.append(["Einstellung", "Wert", "Hinweis"])
    for c in ws[1]:
        c.font, c.fill, c.border = BOLD, HEAD_FILL, HEAD_BORDER
    header_check = ",".join(
        f'data_input!{get_column_letter(i + 1)}1="{name}"'
        for i, name in enumerate(INPUT_COLUMNS))
    dep_check = ",".join(
        f'data_deposits!{get_column_letter(i + 1)}1="{DEPOSIT_COLUMNS[i]}"'
        for i in (2, 7, 14, 15, 16))
    settings = [
        ("Redash CSV-URL KM-Daten", None,
         "CSV-Results-URL der Redash-Query inkl. api_key. Siehe Anleitung."),
        ("Redash CSV-URL Kautionen (optional)", None,
         "Optional. Ohne URL wird keine Kaution verrechnet."),
        ("MwSt-Faktor", 1 if vat is None else vat,
         "1 = Preise inkl. MwSt (Privatkunden). 1.081 = Preise exkl. MwSt "
         "(Firmenkunden). Gilt für alle CHF-Preise in web und Übersicht, "
         "nicht für die Kaution."),
        ("Stichtag (leer = heute)", as_of,
         "Leer lassen = heute. Nur für Tests oder Rückblicke ausfüllen."),
        ("Stichtag effektiv", '=IF(LEN(B5)=0,TODAY(),IF(ISNUMBER(B5),'
         'ROUND(B5-0.5,0),IFERROR(DATE(VALUE(LEFT(TRIM(B5),4)),'
         'VALUE(MID(TRIM(B5),6,2)),VALUE(MID(TRIM(B5),9,2))),TODAY())))',
         "Wird berechnet. Nicht überschreiben."),
        ("MwSt-Faktor effektiv", '=IF(LEN(B4)=0,1,IFERROR(IF(AND(N(B4)>=1,'
         'N(B4)<=1.2),N(B4),IF(AND(VALUE(TRIM(B4))>=1,VALUE(TRIM(B4))<=1.2),'
         'VALUE(TRIM(B4)),NA())),NA()))',
         "Wird berechnet. Leer = 1. Erlaubt sind 1 bis 1.2. Sonst steht hier "
         "#N/A und alle CHF-Preise in web und Übersicht bleiben leer."),
        ("Tarifwechsel Extra-KM", tariffs["cutoff"],
         "Abo-Start vor diesem Datum = alter Tarif (nach Monatspreis), ab "
         "diesem Datum = neuer Tarif (nach Listenpreis). Abo-Start = erste "
         "Übergabe der Fahrzeug-Kette."),
        ("Import-Check KM-Daten",
         f'=IF(NOT({LOCALE_OK}),"{LOCALE_MSG}",IF(ISERROR(data_input!A1),'
         'IF(LEN(B2)=0,"Trag zuerst die Redash-URL in B2 ein.",'
         '"Achtung: Import fehlgeschlagen. Schau dir den Fehler in '
         'data_input!A1 an (URL / api_key prüfen)."),'
         f'IF(D9=1,IF(LEN(data_input!$B${LAST + 1})>0,'
         f'"Achtung: mehr als {LAST - 1} Zeilen. Nur die ersten {LAST - 1} '
         'werden berechnet.","OK"),IF(AND(LEN(B2)=0,LEN(data_input!B1)=0),'
         '"Trag zuerst die Redash-URL in B2 ein.",'
         '"Achtung: Spalten in data_input passen nicht zur Vorlage. '
         'Bis dahin bleibt «web» leer."))))',
         "Erwartet die Spalten vehicle … monthly_price (A–T) in genau dieser "
         "Reihenfolge. Passt was nicht (D9 = 0), bleibt «web» leer."),
        ("Import-Check Kautionen",
         f'=IF(NOT({LOCALE_OK}),"{LOCALE_MSG}",IF(ISERROR(data_deposits!A1),'
         'IF(LEN(B3)=0,"Keine Kautions-URL (optional).",'
         '"Achtung: Kautions-Import fehlgeschlagen. Schau dir den Fehler in '
         'data_deposits!A1 an (URL / api_key prüfen)."),'
         f'IF(D10=1,IF(LEN(data_deposits!$Q${DEP_LAST + 1})>0,'
         f'"Achtung: mehr als {DEP_LAST - 1} Zeilen. Nur die ersten '
         f'{DEP_LAST - 1} werden berechnet.","OK"),'
         'IF(AND(LEN(B3)=0,LEN(data_deposits!C1)=0),'
         '"Keine Kautions-URL (optional).",'
         '"Achtung: Spalten in data_deposits passen nicht zur Vorlage."))))',
         "Erwartet amount in C, subscription_status in H, status in O, "
         "reason in P, license_plate_number in Q. Ist die URL gesetzt und "
         "der Import kaputt (D10 = 0), bleiben Kaution und Mehrkilometer "
         "nach Kaution leer."),
    ]
    for label, value, hint in settings:
        ws.append([label, value, hint])
    # Interne Schalter: 1 = Import brauchbar. calc rechnet nur dann.
    ws["D1"] = "Check (1 = ok)"
    ws["D1"].font, ws["D1"].fill, ws["D1"].border = (BOLD, HEAD_FILL,
                                                     HEAD_BORDER)
    ws["D9"] = f'=IFERROR(IF(AND({LOCALE_OK},{header_check}),1,0),0)'
    ws["D10"] = f'=IFERROR(IF(AND({LOCALE_OK},{dep_check}),1,0),0)'
    for ref in ("D9", "D10"):
        ws[ref].number_format = "0"
        ws[ref].fill = PatternFill("solid", fgColor="F3F3F3")
    ws["B4"].number_format = "0.000"
    ws["B5"].number_format = FMT_DATE
    ws["B6"].number_format = FMT_DATE
    ws["B7"].number_format = "0.000"
    ws["B8"].number_format = FMT_DATE
    for r in range(2, 11):
        ws.cell(r, 1).font = BOLD
        ws.cell(r, 3).font = NOTE_FONT
        ws.cell(r, 3).alignment = WRAP
    for r in (6, 7, 9, 10):
        ws.cell(r, 2).fill = PatternFill("solid", fgColor="F3F3F3")

    def table(title_row, title, head, rows, note_last):
        ws.cell(title_row, 1, title).font = Font(bold=True, size=12)
        for i, h in enumerate(head, 1):
            c = ws.cell(title_row + 1, i, h)
            c.font, c.fill, c.border = BOLD, HEAD_FILL, HEAD_BORDER
        for i, (upper, price) in enumerate(rows):
            r = title_row + 2 + i
            ws.cell(r, 1, "darüber" if upper is None else upper).number_format = FMT_KM
            ws.cell(r, 2, price).number_format = "0.00"
        ws.cell(title_row + 1 + len(rows), 3, note_last).font = NOTE_FONT

    table(OLD_ROW - 2,
          "Extra-KM-Tarif alt (Abo-Start vor B8): nach Monatspreis",
          ["Monatspreis bis CHF (inkl.)", "Preis pro km CHF (inkl. MwSt)",
           "Hinweis"], tariffs["old"],
          "Leer = kein Tarif. Dann wird kein Extra-KM-Preis angezeigt.")
    table(NEW_ROW - 2,
          "Extra-KM-Tarif neu (Abo-Start ab B8): nach Listenpreis",
          ["Listenpreis bis CHF (inkl.)", "Preis pro km CHF (inkl. MwSt)",
           "Hinweis"], tariffs["new"], "Alles über der letzten Grenze.")

    c = ws.cell(PKG_ROW - 1, 1, "Verfügbare KM-Pakete (km / Monat)")
    c.font, c.fill, c.border = BOLD, HEAD_FILL, HEAD_BORDER
    ws.cell(PKG_ROW - 1, 3, "Für die Paket-Empfehlung. Freie Zeilen für "
                           "weitere Pakete.").font = NOTE_FONT
    for i, p in enumerate(tariffs["packages"][:PACKAGE_SLOTS]):
        ws.cell(PKG_ROW + i, 1, p).number_format = FMT_KM

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 80
    ws.column_dimensions["D"].width = 14
    ws.freeze_panes = "A2"


def build_calc(ws) -> None:
    for key, header, _expr, _fmt in CALC:
        ws[f"{COL[key]}1"] = header
    for key, header, _expr in DEPOSIT_HELPER:
        ws[f"{COL[key]}1"] = header
    for c in ws[1]:
        if c.value:
            c.font, c.fill, c.alignment = BOLD, HEAD_FILL, CENTER_WRAP
    ws.row_dimensions[1].height = 60

    for r in range(FIRST, LAST + 1):
        for key, _header, expr, fmt in CALC:
            c = ws[f"{COL[key]}{r}"]
            c.value = calc_formula(key, expr, r)
            if fmt:
                c.number_format = fmt
    for r in range(FIRST, DEP_LAST + 1):
        for key, _header, expr in DEPOSIT_HELPER:
            body = expr.format_map(_Refs(r))
            ws[f"{COL[key]}{r}"] = "=" + body

    for key in COL:
        ws.column_dimensions[COL[key]].width = 15
    ws.column_dimensions[COL["vehicle"]].width = 26
    ws.column_dimensions[COL["chain_key"]].width = 22
    ws.freeze_panes = f"{COL['vehicle']}2"


def web_formula(key: str, kind: str, rank: int) -> str:
    v = pick(key, rank)
    if kind == "text":
        return f'=IFERROR({v}&"","")'
    if kind == "vat":
        v = f"{v}/{CFG_VAT}"
    elif kind == "rp":
        v = f"{v}*100/{CFG_VAT}"
    return f'=IFERROR(ROUND({v},0),"")'


def build_web(ws) -> None:
    ws.append(CONTRACT)
    for c in ws[1]:
        c.font = BOLD
    for rank in range(1, WEB_ROWS + 1):
        row = rank + 1
        for i, name in enumerate(CONTRACT, 1):
            key, kind = WEB[name]
            c = ws.cell(row, i, web_formula(key, kind, rank))
            if kind != "text":
                c.number_format = FMT_INT
    for i in range(1, len(CONTRACT) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.freeze_panes = "A2"


def overview_formula(key: str, kind: str, rank: int) -> str:
    v = pick(key, rank)
    if kind == "vat":
        v = f"{v}/{CFG_VAT}"
    return f'=IFERROR({v},"")'


def build_overview(ws) -> None:
    ncol = len(OVERVIEW)
    ws["A1"] = "CARIFY KM-Tracker"
    ws["A1"].font = TITLE
    kpis = [
        ("D1", "Stichtag", "E1", f"={CFG_ASOF}", FMT_DATE),
        ("G1", "Aktive Abos", "H1", f'=COUNTIF({calc_ref("active")},1)', "0"),
        ("J1", "Preise", "K1", f'=IF(ISERROR({CFG_VAT}),"MwSt-Faktor '
         f'ungültig (Config!B4: 1 oder 1.081)",IF({CFG_VAT}=1,"inkl. MwSt",'
         '"exkl. MwSt"))', None),
        ("M1", "Import", "N1", '=IF(Config!$B$9<>"OK",Config!$B$9,'
         'IF(OR(Config!$B$10="OK",LEN(Config!$B$3)=0),"OK",'
         '"Kautionen: "&Config!$B$10))', None),
    ]
    for lc, label, vc, formula, fmt in kpis:
        ws[lc] = label
        ws[lc].font = NOTE_FONT
        ws[lc].alignment = Alignment(horizontal="right")
        ws[vc] = formula
        ws[vc].font = BOLD
        if fmt:
            ws[vc].number_format = fmt
    ws.row_dimensions[1].height = 28

    # grouped header (row 2) + column header (row 3)
    start = 1
    for i in range(1, ncol + 1):
        group = OVERVIEW[i - 1][0]
        nxt = OVERVIEW[i][0] if i < ncol else None
        if nxt != group:
            ws.cell(2, start, group)
            if i > start:
                ws.merge_cells(start_row=2, start_column=start, end_row=2,
                               end_column=i)
            start = i + 1
    for i, (_g, header, _k, _kind, _fmt, width) in enumerate(OVERVIEW, 1):
        g = ws.cell(2, i)
        g.font, g.fill, g.alignment = BOLD, GROUP_FILL, CENTER_WRAP
        h = ws.cell(3, i, header)
        h.font, h.fill, h.alignment, h.border = (BOLD, HEAD_FILL, CENTER_WRAP,
                                                 HEAD_BORDER)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[3].height = 45

    for rank in range(1, OVERVIEW_ROWS + 1):
        row = rank + 3
        for i, (_g, _h, key, kind, fmt, _w) in enumerate(OVERVIEW, 1):
            c = ws.cell(row, i, overview_formula(key, kind, rank))
            if fmt:
                c.number_format = fmt
            if key == "insurance":
                c.alignment = Alignment(horizontal="center")

    # rot = über dem Paket, orange = Prognose über dem Paket
    last = 3 + OVERVIEW_ROWS
    red = PatternFill("solid", fgColor="F4CCCC")
    orange = PatternFill("solid", fgColor="FCE5CD")
    for header, fill in (("Differenz", red), ("Mehrkilometer nach Kaution", red),
                         ("Prognose Differenz Ende Mindestlaufzeit", orange)):
        idx = [o[1] for o in OVERVIEW].index(header) + 1
        col = get_column_letter(idx)
        ws.conditional_formatting.add(
            f"{col}4:{col}{last}",
            FormulaRule(formula=[f"AND(ISNUMBER({col}4),{col}4>0)"], fill=fill))
    ws.freeze_panes = "C4"


def build_import(ws, url_cell: str, fallback: str) -> None:
    # Kein IFERROR: echte Fehler (api_key, Redash down) bleiben in A1
    # sichtbar, Config!B9/B10 melden sie.
    ws["A1"] = f'=IF(LEN({url_cell})=0,"{fallback}",IMPORTDATA({url_cell}))'


# --------------------------------------------------------------------------
# Sample mode (tests only)
# --------------------------------------------------------------------------
_INT = re.compile(r"^-?\d+$")
_ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def sample_value(text: str):
    """Mimic IMPORTDATA typing: plain integers and pure ISO dates become
    numbers/dates, everything else (e.g. "4100.0", "...T00:00:00") stays text."""
    if text == "":
        return None
    if _INT.match(text):
        return int(text)
    if _ISO.match(text):
        return dt.date.fromisoformat(text)
    return text


def write_sample(ws, path: str) -> None:
    with open(path, newline="", encoding="utf-8") as fh:
        for i, row in enumerate(csv.reader(fh)):
            ws.append(row if i == 0 else [sample_value(v) for v in row])


# --------------------------------------------------------------------------
def build(out: str, sample: str | None = None, sample_deposits: str | None = None,
          as_of: dt.date | None = None, vat: float | None = None,
          tariffs: str | None = None) -> None:
    tariff_data = load_tariffs(tariffs)
    wb = Workbook()
    names = ["Anleitung", "Übersicht", "web", "calc", "data_input",
             "data_deposits", "Config"]
    ws = {n: (wb.active if i == 0 else wb.create_sheet()) for i, n in
          enumerate(names)}
    for n, sheet in ws.items():
        sheet.title = n
    colors = {"Anleitung": "6AA84F", "Übersicht": "3C78D8", "web": "E69138",
              "calc": "999999", "data_input": "CCCCCC",
              "data_deposits": "CCCCCC", "Config": "F1C232"}
    for n, color in colors.items():
        ws[n].sheet_properties.tabColor = color

    build_anleitung(ws["Anleitung"])
    build_overview(ws["Übersicht"])
    build_web(ws["web"])
    build_calc(ws["calc"])
    if sample:
        write_sample(ws["data_input"], sample)
    else:
        build_import(ws["data_input"], "Config!$B$2",
                     "Bitte Redash-URL in Config!B2 eintragen")
    if sample_deposits:
        write_sample(ws["data_deposits"], sample_deposits)
    elif not sample:
        build_import(ws["data_deposits"], "Config!$B$3", "")
    build_config(ws["Config"], vat, as_of if sample else None, tariff_data)
    wb.active = 1
    wb.save(out)


def main() -> None:
    global LAST, DEP_LAST, OVERVIEW_ROWS
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("--out", required=True, help="output .xlsx path")
    p.add_argument("--sample", help="fixture CSV for data_input (tests only)")
    p.add_argument("--sample-deposits",
                   help="fixture CSV for data_deposits (with --sample)")
    p.add_argument("--as-of", type=dt.date.fromisoformat,
                   help="Stichtag YYYY-MM-DD in Config!B5 (with --sample)")
    p.add_argument("--vat", type=float, help="MwSt-Faktor in Config!B4 "
                                             "(default 1)")
    p.add_argument("--tariffs", help="Tarif-JSON (default: tariffs.local.json "
                                     "neben diesem Skript)")
    p.add_argument("--rows", type=int, default=LAST - 1,
                   help="max. KM-Datenzeilen (default %(default)s); kleiner = "
                        "kleinere Datei")
    p.add_argument("--deposit-rows", type=int, default=DEP_LAST - 1,
                   help="max. Kautions-Datenzeilen (default %(default)s)")
    a = p.parse_args()
    if (a.sample_deposits or a.as_of) and not a.sample:
        p.error("--sample-deposits/--as-of only work together with --sample")
    LAST, DEP_LAST = FIRST + a.rows - 1, FIRST + a.deposit_rows - 1
    OVERVIEW_ROWS = min(OVERVIEW_ROWS, a.rows)
    build(a.out, a.sample, a.sample_deposits, a.as_of, a.vat, a.tariffs)


if __name__ == "__main__":
    main()
